"""Đối soát file định import vào Power BI với Gold local và lưu bằng chứng."""

from __future__ import annotations

import hashlib
import json
from decimal import Decimal
from pathlib import Path

import duckdb
import matplotlib
from openpyxl import load_workbook

matplotlib.use("Agg")

import matplotlib.pyplot as plt

ROOT = Path(__file__).resolve().parents[1]
DB_PATH = ROOT / "data/tmp/week4/warehouse.duckdb"
EXPORT_DIR = ROOT / "data/exports/powerbi"
REPORT_DIR = ROOT / "reports/week10"


def sha256(path: Path) -> str:
    return hashlib.sha256(path.read_bytes()).hexdigest()


def verify(db_path: Path, manifest_path: Path) -> dict:
    manifest = json.loads(manifest_path.read_text(encoding="utf-8"))
    workbook_path = manifest_path.parent / manifest["file"]
    if sha256(workbook_path) != manifest["file_sha256"]:
        raise ValueError("Workbook checksum không khớp manifest")
    if sha256(db_path) != manifest["warehouse_sha256"]:
        raise ValueError("Gold đã đổi kể từ lúc export — chạy lại make week10")

    workbook = load_workbook(workbook_path, read_only=True, data_only=True)
    try:
        expected_tables = set(manifest["sheets"])
        if set(workbook.sheetnames) != expected_tables:
            raise ValueError("Danh sách sheet không khớp manifest")
        for sheet_name, expected_rows in manifest["sheets"].items():
            if workbook[sheet_name].max_row - 1 != expected_rows:
                raise ValueError(f"Số dòng {sheet_name} không khớp manifest")

        rows = list(workbook["FactSpend"].values)
        headers = rows[0]
        fact = [dict(zip(headers, row, strict=True)) for row in rows[1:]]
        spend = sum((Decimal(str(row["amount_usd"])) for row in fact), Decimal())
        advertisers = len({row["advertiser_id"] for row in fact})
        weeks = sorted({row["week_start_date"].date() for row in fact})
        latest = weeks[-1]
        latest_rows = [row for row in fact if row["week_start_date"].date() == latest]
        latest_spend = sum(
            (Decimal(str(row["amount_usd"])) for row in latest_rows), Decimal()
        )
        latest_top10 = sum(
            sorted(
                (Decimal(str(row["amount_usd"])) for row in latest_rows),
                reverse=True,
            )[:10],
            Decimal(),
        )
    finally:
        workbook.close()

    con = duckdb.connect(str(db_path), read_only=True)
    try:
        sql_spend, sql_rows, sql_advertisers = con.execute("""
            SELECT SUM(amount_usd), COUNT(*), COUNT(DISTINCT advertiser_id)
            FROM fct_advertiser_weekly_spend
        """).fetchone()
        sql_latest_spend, sql_latest_advertisers = con.execute(
            """
            SELECT total_spend_usd, advertiser_count
            FROM mart_bi_weekly WHERE week_start_date = ?
        """,
            [latest],
        ).fetchone()
        sql_top10 = con.execute(
            """
            SELECT SUM(amount_usd) FROM (
                SELECT amount_usd FROM fct_advertiser_weekly_spend
                WHERE week_start_date = ?
                ORDER BY amount_usd DESC, advertiser_id LIMIT 10
            )
        """,
            [latest],
        ).fetchone()[0]
        dq = con.execute(
            """
            SELECT COUNT(*) FILTER (WHERE status = 'PASS'), COUNT(*)
            FROM audit.dq_scorecard WHERE run_id = ?
        """,
            [manifest["latest_audit_run_id"]],
        ).fetchone()
    finally:
        con.close()

    checks = {
        "fact_rows": (len(fact), sql_rows),
        "total_spend_usd": (spend, sql_spend),
        "distinct_advertisers": (advertisers, sql_advertisers),
        "latest_week_spend_usd": (latest_spend, sql_latest_spend),
        "latest_week_advertisers": (len(latest_rows), sql_latest_advertisers),
        "latest_week_top10_spend_usd": (latest_top10, sql_top10),
    }
    for metric, (excel_value, sql_value) in checks.items():
        if Decimal(str(excel_value)) != Decimal(str(sql_value)):
            raise ValueError(
                f"Đối soát thất bại: {metric}: {excel_value} != {sql_value}"
            )
    if dq[0] != dq[1] or not dq[1]:
        raise ValueError("DQ của audit run mới nhất chưa PASS")

    return {
        "status": "PASS_LOCAL_ONLY",
        "power_bi_status": "NOT_VERIFIED",
        "refresh_mode": "STATIC_SNAPSHOT",
        "workbook": workbook_path.name,
        "workbook_sha256": manifest["file_sha256"],
        "warehouse_sha256": manifest["warehouse_sha256"],
        "period": [weeks[0].isoformat(), latest.isoformat()],
        "dq_passed": dq[0],
        "dq_total": dq[1],
        "checks": {
            metric: {"excel": str(pair[0]), "sql": str(pair[1]), "match": True}
            for metric, pair in checks.items()
        },
        "weekly_spend": [
            {
                "week_start": week.isoformat(),
                "spend_usd": str(
                    sum(
                        (
                            Decimal(str(row["amount_usd"]))
                            for row in fact
                            if row["week_start_date"].date() == week
                        ),
                        Decimal(),
                    )
                ),
            }
            for week in weeks
        ],
    }


def plot_preview(evidence: dict, path: Path) -> None:
    weekly = evidence["weekly_spend"]
    dates = [row["week_start"][5:] for row in weekly]
    amounts = [float(row["spend_usd"]) / 1_000_000 for row in weekly]
    fig, ax = plt.subplots(figsize=(9, 4.5))
    bars = ax.bar(dates, amounts, color="#255c9a")
    ax.bar_label(bars, fmt="%.1f")
    ax.set(title="BI source preview — observed spend", ylabel="USD millions")
    ax.text(
        0,
        -0.18,
        "Local Gold snapshot · 4 loaded weeks · Not a Power BI screenshot",
        transform=ax.transAxes,
        fontsize=9,
        color="#555",
    )
    ax.grid(axis="y", alpha=0.2)
    ax.set_axisbelow(True)
    fig.savefig(path, dpi=180, bbox_inches="tight", facecolor="white")
    plt.close(fig)


def main() -> None:
    evidence = verify(DB_PATH, EXPORT_DIR / "snapshot_manifest.json")
    REPORT_DIR.mkdir(parents=True, exist_ok=True)
    chart_dir = REPORT_DIR / "charts"
    chart_dir.mkdir(exist_ok=True)
    (REPORT_DIR / "evidence.json").write_text(
        json.dumps(evidence, indent=2, ensure_ascii=False) + "\n",
        encoding="utf-8",
    )
    plot_preview(evidence, chart_dir / "01_local_bi_source.png")
    print(json.dumps(evidence["checks"], indent=2))


if __name__ == "__main__":
    main()
