from __future__ import annotations

import hashlib
import json
from datetime import date, datetime, timezone
from decimal import Decimal
from pathlib import Path

import duckdb
from openpyxl import Workbook
from openpyxl.styles import Font
from openpyxl.utils import get_column_letter
from openpyxl.worksheet.table import Table, TableStyleInfo

ROOT = Path(__file__).resolve().parents[1]
DB_PATH = ROOT / "data/tmp/week4/warehouse.duckdb"
OUTPUT_DIR = ROOT / "data/exports/powerbi"

QUERIES = {
    "FactSpend": """
        SELECT
            advertiser_week_key,
            advertiser_id,
            week_start_date,
            amount_usd,
            currency_code
        FROM fct_advertiser_weekly_spend
        ORDER BY week_start_date, advertiser_id
    """,
    "DimAdvertiser": """
        SELECT advertiser_id, advertiser_name, last_observed_week
        FROM dim_advertiser
        ORDER BY advertiser_id
    """,
    "DimWeek": """
        SELECT
            week_start_date,
            CAST(week_end_date AS DATE) AS week_end_date,
            calendar_year,
            calendar_month
        FROM dim_week
        ORDER BY week_start_date
    """,
    "DQScorecard": """
        SELECT
            run_id, check_name, status, severity, observed, expected
        FROM audit.dq_scorecard
        WHERE run_id = (
            SELECT run_id
            FROM audit.pipeline_runs
            ORDER BY finished_at_utc DESC
            LIMIT 1
        )
        ORDER BY check_name
    """,
    "TargetingStatus": """
        SELECT
            required_source,
            data_status,
            targeting_metrics_ready,
            source_checked_on,
            latest_loaded_week,
            reason
        FROM mart_targeting_coverage
        ORDER BY required_source
    """,
}


def excel_value(value):
    if isinstance(value, Decimal):
        # Excel lưu số; trong Power BI đặt lại là Fixed decimal number
        return float(value)
    return value


def add_table(workbook, name: str, columns: list[str], rows: list[tuple]) -> None:
    sheet = workbook.create_sheet(name)
    sheet.append(columns)

    for row in rows:
        values = [excel_value(value) for value in row]
        sheet.append(values)

        # Giữ ngày dưới dạng Excel date thay vì chuỗi văn bản
        row_number = sheet.max_row
        for column_number, value in enumerate(values, start=1):
            if isinstance(value, date):
                sheet.cell(row_number, column_number).number_format = "yyyy-mm-dd"

    for cell in sheet[1]:
        cell.font = Font(bold=True)

    sheet.freeze_panes = "A2"

    table = Table(
        displayName=name,
        ref=f"A1:{get_column_letter(len(columns))}{sheet.max_row}",
    )
    table.tableStyleInfo = TableStyleInfo(
        name="TableStyleMedium2",
        showRowStripes=True,
    )
    sheet.add_table(table)


def main() -> None:
    if not DB_PATH.is_file():
        raise SystemExit("Chưa có Gold local — kiểm tra warehouse hoặc chạy make week6")

    OUTPUT_DIR.mkdir(parents=True, exist_ok=True)

    connection = duckdb.connect(str(DB_PATH), read_only=True)
    connection.execute("SET memory_limit = '1GB'")
    connection.execute("SET threads = 2")

    try:
        first_week, last_week, fact_rows, advertisers, total_usd = (
            connection.execute("""
                SELECT
                    MIN(week_start_date),
                    MAX(week_start_date),
                    COUNT(*),
                    COUNT(DISTINCT advertiser_id),
                    SUM(amount_usd)
                FROM fct_advertiser_weekly_spend
            """).fetchone()
        )

        latest_audit = connection.execute("""
            SELECT
                run_id,
                status,
                CAST(finished_at_utc AS VARCHAR)
            FROM audit.pipeline_runs
            ORDER BY finished_at_utc DESC
            LIMIT 1
        """).fetchone()

        workbook = Workbook()
        workbook.remove(workbook.active)

        row_counts = {}

        for name, sql in QUERIES.items():
            result = connection.execute(sql)
            columns = [description[0] for description in result.description]
            rows = result.fetchall()
            add_table(workbook, name, columns, rows)
            row_counts[name] = len(rows)
            print(f"{name}: {len(rows)} rows")

        exported_at = datetime.now(timezone.utc).isoformat(timespec="seconds")
        warehouse_hash = hashlib.sha256(DB_PATH.read_bytes()).hexdigest()

        add_table(
            workbook,
            "SnapshotInfo",
            [
                "refresh_mode",
                "first_week",
                "last_week",
                "exported_at_utc",
                "warehouse_sha256",
                "latest_audit_run_id",
                "latest_audit_status",
                "scope",
            ],
            [
                (
                    "STATIC_SNAPSHOT",
                    first_week,
                    last_week,
                    exported_at,
                    warehouse_hash,
                    latest_audit[0],
                    latest_audit[1],
                    "Observed USD spend in loaded weeks; not US-only",
                )
            ],
        )
        row_counts["SnapshotInfo"] = 1

        output = OUTPUT_DIR / f"political_ads_gold_{last_week}.xlsx"
        workbook.save(output)

        manifest = {
            "file": output.name,
            "file_sha256": hashlib.sha256(output.read_bytes()).hexdigest(),
            "warehouse_sha256": warehouse_hash,
            "exported_at_utc": exported_at,
            "first_week": first_week.isoformat(),
            "last_week": last_week.isoformat(),
            "fact_rows": fact_rows,
            "distinct_advertisers": advertisers,
            "total_spend_usd": str(total_usd),
            "latest_audit_run_id": latest_audit[0],
            "latest_audit_status": latest_audit[1],
            "sheets": row_counts,
            "refresh_mode": "STATIC_SNAPSHOT",
        }

        manifest_path = OUTPUT_DIR / "snapshot_manifest.json"
        manifest_path.write_text(
            json.dumps(manifest, indent=2, ensure_ascii=False) + "\n",
            encoding="utf-8",
        )

        print(f"\nWorkbook: {output}")
        print(f"Manifest: {manifest_path}")
        print(f"Spend USD: {total_usd}")
    finally:
        connection.close()


if __name__ == "__main__":
    main()
