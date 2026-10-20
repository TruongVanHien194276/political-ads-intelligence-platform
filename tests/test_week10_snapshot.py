import json
from datetime import date

import pytest
from openpyxl import Workbook

from scripts.week10_export_powerbi import add_table, excel_value
from scripts.week10_verify_snapshot import verify


def test_excel_table_keeps_dates_and_numbers():
    workbook = Workbook()
    add_table(
        workbook,
        "FactSpend",
        ["week_start_date", "amount_usd"],
        [(date(2026, 10, 11), 123)],
    )
    sheet = workbook["FactSpend"]
    assert sheet.tables["FactSpend"].ref == "A1:B2"
    assert sheet["A2"].value == date(2026, 10, 11)
    assert sheet["A2"].number_format == "yyyy-mm-dd"
    assert excel_value(123) == 123


def test_verifier_rejects_changed_workbook(tmp_path):
    workbook = tmp_path / "snapshot.xlsx"
    workbook.write_bytes(b"changed")
    manifest = tmp_path / "snapshot_manifest.json"
    manifest.write_text(
        json.dumps({"file": workbook.name, "file_sha256": "wrong"}),
        encoding="utf-8",
    )
    with pytest.raises(ValueError, match="checksum"):
        verify(tmp_path / "warehouse.duckdb", manifest)
